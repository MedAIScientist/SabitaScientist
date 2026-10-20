"""The study dossier: one Word document that says what the numbers rest on.

It gathers the provenance (which data version, which rows, which engine), the
participant flow, the data handling, the missing-data treatment and the results the
researcher chose — each result with the population fingerprint it was computed on.

Two rules this module keeps:

* it takes plain data and returns bytes: no database, no FastAPI, so the document can
  be produced by a worker later;
* it is not a compliance certificate. The first paragraph after the title says so, and
  no wording anywhere claims conformance with a reporting guideline.
"""

from __future__ import annotations

import io
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

__all__ = ["build_dossier"]

_DISCLAIMER = (
    "This dossier gathers the records of one analysis. It is not a compliance "
    "certificate and not an electronically signed record: it does not certify that the "
    "study conforms to STROBE, CONSORT, PRISMA or any other reporting guideline. It "
    "records what was computed, on which rows, from which data version."
)

_MISSING_NO_ASSESSMENT = (
    "No missing-data assessment was included in this dossier. Little's MCAR test and "
    "multiple imputation are available; a study reporting complete-case analyses only "
    "should say so explicitly in its methods."
)


def _add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths: list[float] | None = None,
):
    """A bordered table; ``rows`` may be empty, in which case only the header is written."""
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = ""
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.size = Pt(9)
    for values in rows:
        cells = table.add_row().cells
        for index, text in enumerate(values):
            cells[index].text = ""
            run = cells[index].paragraphs[0].add_run(text)
            run.font.size = Pt(9)
    if widths:
        for row in table.rows:
            for index, width in enumerate(widths):
                if index < len(row.cells):
                    row.cells[index].width = width
    return table


def _add_note(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(8)


def _bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def _section_provenance(
    doc: Document,
    table: dict[str, Any],
    project: dict[str, Any],
    versions: list[dict[str, Any]],
    population: dict[str, Any] | None,
) -> None:
    doc.add_heading("1. Study and data provenance", level=1)
    origin = str(table.get("data_origin") or "public")
    lines = [
        f"Project: {project.get('name') or project.get('id')}",
        f"Cohort table: {table.get('name')}",
        f"Data origin: {origin}"
        + (
            f" (governed dataset {table.get('dataset_id')})"
            if table.get("dataset_id")
            else ""
        ),
        f"Table created: {table.get('created_at')}",
    ]
    _bullets(doc, lines)
    if origin != "governed":
        _add_note(
            doc,
            "Public or synthetic data: no ethics approval is linked to this cohort, "
            "because none is required for it.",
        )

    doc.add_heading("Saved versions", level=2)
    _add_table(
        doc,
        ["Version", "Rows", "Columns", "Content (sha256)", "Created", "Note"],
        [
            [
                f"v{version.get('version_no')}",
                f"{int(version.get('rows') or 0):,}",
                str(version.get("columns") or ""),
                str(version.get("content_sha256") or "")[:16],
                str(version.get("created_at") or ""),
                str(version.get("note") or ""),
            ]
            for version in versions
        ],
        widths=None,
    )

    doc.add_heading("Analysis population", level=2)
    if not population:
        doc.add_paragraph("No analysis population was declared.")
        return
    doc.add_paragraph(
        f"{population.get('name')}: {int(population.get('n_included') or 0):,} of "
        f"{int(population.get('n_base') or 0):,} rows."
    )
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(
        f"Fingerprint of the analysed rows: {population.get('fingerprint')}"
    )
    run.font.size = Pt(9)
    criteria = population.get("criteria") or []
    if criteria:
        _add_table(
            doc,
            ["Criterion", "Family", "Kind"],
            [
                [
                    str(item.get("name") or ""),
                    str(item.get("family") or ""),
                    str(item.get("kind") or ""),
                ]
                for item in criteria
            ],
        )
    counts = population.get("counts") or []
    if counts:
        doc.add_heading("Rows each criterion accounts for", level=2)
        _add_table(
            doc,
            ["Criterion", "Fires on", "First reason for"],
            [
                [
                    str(item.get("name") or ""),
                    f"{int(item.get('marginal_base') or 0):,}",
                    f"{int(item.get('sequential') or 0):,}",
                ]
                for item in counts
            ],
        )


def _section_flow(doc: Document, flow: dict[str, Any] | None) -> None:
    doc.add_heading("2. Participant flow (STROBE item 13)", level=1)
    if not flow:
        doc.add_paragraph("No participant flow was recorded.")
        return
    rows: list[list[str]] = []
    for box in flow.get("main") or []:
        exits = [
            item
            for item in flow.get("exits") or []
            if item.get("after") == box.get("id")
        ]
        reasons = "; ".join(
            f"{item.get('label')} ({int(item.get('n') or 0):,})" for item in exits
        )
        rows.append(
            [
                str(box.get("label") or ""),
                f"{int(box.get('n') or 0):,}",
                reasons or "—",
                "; ".join(box.get("details") or []) or "",
            ]
        )
    for box in flow.get("analyses") or []:
        rows.append(
            [str(box.get("label") or ""), f"{int(box.get('n') or 0):,}", "—", ""]
        )
    _add_table(doc, ["Step", "Rows", "Excluded", "Detail"], rows)
    if flow.get("inconsistencies"):
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(
            "The flow does not add up and the trace is therefore incomplete: "
            + "; ".join(flow["inconsistencies"])
        )
        run.bold = True
        run.font.size = Pt(9)
    if flow.get("notes"):
        _bullets(doc, [str(note) for note in flow["notes"]])


def _section_handling(
    doc: Document, steps: list[dict[str, Any]], quality: dict[str, Any] | None
) -> None:
    doc.add_heading("3. Data handling", level=1)
    changes = [
        step
        for step in steps
        if step.get("kind") in ("derive", "filter", "merge", "impute")
    ]
    if changes:
        _add_table(
            doc,
            ["Step", "Kind", "Detail", "Rows in", "Rows out"],
            [
                [
                    str(step.get("seq")),
                    str(step.get("kind") or ""),
                    str(
                        (step.get("params") or {}).get("name")
                        or (step.get("params") or {}).get("action")
                        or ""
                    ),
                    str((step.get("data") or {}).get("rows_in") or ""),
                    str((step.get("data") or {}).get("rows_out") or ""),
                ]
                for step in changes
            ],
        )
        _add_note(
            doc,
            "Every change of the rows is one link of a hash-chained trace; the app can "
            "replay them and report any step that no longer reproduces the recorded table.",
        )
    else:
        doc.add_paragraph(
            "The rows were analysed as imported: no derived variable and no row subset "
            "was applied."
        )

    if quality:
        doc.add_heading("Data quality of the analysed version", level=2)
        doc.add_paragraph(
            f"Score {float(quality.get('score') or 0):.0%} ({quality.get('band')}), on "
            f"{int(quality.get('rows') or 0):,} rows."
        )
        _add_table(
            doc,
            ["Dimension", "Score", "Band"],
            [
                [
                    str(item.get("label") or ""),
                    f"{float(item.get('score') or 0):.0%}",
                    str(item.get("band") or ""),
                ]
                for item in quality.get("dimensions") or []
            ],
        )
        for warning in quality.get("warnings") or []:
            _add_note(doc, str(warning))


def _section_missing(doc: Document, analyses: list[dict[str, Any]]) -> None:
    doc.add_heading("4. Missing data", level=1)
    mcar = [item for item in analyses if item.get("method") == "mcar"]
    pooled = [item for item in analyses if item.get("method") == "multiple_imputation"]
    if not mcar and not pooled:
        doc.add_paragraph(_MISSING_NO_ASSESSMENT)
    for analysis in mcar:
        result = analysis.get("result") or {}
        doc.add_heading("Little's MCAR test", level=2)
        doc.add_paragraph(
            f"chi-square = {float(result.get('statistic') or 0):.3f}, df = "
            f"{int(result.get('df') or 0)}, p = {_p(result.get('p_value'))}. "
            + (
                "The data are not consistent with MCAR."
                if result.get("evidence_against_mcar")
                else "No evidence against MCAR — which does not prove it."
            )
        )
        _add_note(
            doc,
            "The test reads the numeric variables only, and has little power in a small "
            "cohort: failing to reject MCAR is not evidence that the data are MCAR.",
        )
    for analysis in pooled:
        result = analysis.get("result") or {}
        doc.add_heading("Multiple imputation", level=2)
        target = {
            "compare": "a two-group comparison",
            "logistic": "a logistic model",
            "cox": "a Cox model",
        }.get(str(result.get("target")), str(result.get("target")))
        imputed = result.get("imputed") or {}
        doc.add_paragraph(
            f"{int(result.get('m') or 0)} imputations, {int(result.get('maxit') or 0)} "
            f"iterations, pooled with Rubin's rules (Barnard-Rubin degrees of freedom), "
            f"for {target}."
        )
        _bullets(
            doc,
            [
                "Imputation model columns: "
                + ", ".join(result.get("model_columns") or []),
                "Imputed cells: "
                + (
                    ", ".join(f"{name} {count}" for name, count in imputed.items())
                    or "none"
                ),
                "Methods: "
                + (
                    ", ".join(
                        f"{name}={method}"
                        for name, method in (result.get("methods") or {}).items()
                    )
                    or ""
                ),
                f"Seed: {result.get('seed')}"
                + (
                    " (the run is reproducible)"
                    if result.get("seed") is not None
                    else ""
                ),
            ],
        )
        terms = result.get("terms") or []
        if terms:
            _bullets(
                doc,
                [
                    f"{term.get('term') or ''!s}: {float(term.get('estimate') or 0):.3f} "
                    f"({float(term.get('ci_low') or 0):.3f} to "
                    f"{float(term.get('ci_high') or 0):.3f}), p {_p(term.get('p_value'))}, "
                    f"FMI {float(term.get('fmi') or 0):.2f}"
                    for term in terms
                ],
            )
            _add_note(doc, "The pooled table is repeated in the results section.")
        worst = max((float(term.get("fmi") or 0) for term in terms), default=0.0)
        if worst > 0.5:
            _add_note(
                doc,
                f"The largest fraction of missing information is {worst:.2f}: at this "
                "level the number of imputations matters, and more than "
                f"{int(result.get('m') or 0)} would be prudent.",
            )
        _add_note(
            doc,
            "Nothing singly imputed was used for inference: a single imputation does not "
            "carry the imputation uncertainty into the standard error.",
        )


def _section_results(doc: Document, analyses: list[dict[str, Any]]) -> None:
    doc.add_heading("5. Results", level=1)
    reported = [
        item for item in analyses if item.get("table") and item.get("method") != "mcar"
    ]
    if not reported:
        doc.add_paragraph("No result was selected for this dossier.")
        return
    for index, analysis in enumerate(reported, start=1):
        doc.add_heading(f"5.{index} {analysis.get('method')}", level=2)
        paragraph = doc.add_paragraph()
        run = paragraph.add_run(
            f"n = {int(analysis.get('n') or 0):,} · population fingerprint "
            f"{analysis.get('population_fingerprint')} · engine {analysis.get('engine_version')} "
            f"· computed {analysis.get('created_at')}"
        )
        run.font.size = Pt(9)
        if analysis.get("stale_reason"):
            stale = doc.add_paragraph()
            run = stale.add_run(
                "Stale result: "
                + str(analysis["stale_reason"])
                + ". It is included because it "
                "was asked for by name, and must not be reported as current."
            )
            run.bold = True
            run.font.size = Pt(9)
        rendered = analysis.get("table") or {}
        columns = rendered.get("columns") or []
        rows = rendered.get("rows") or []
        if columns and isinstance(columns[0], dict):
            header = ["Variable"] + [
                str(column.get("label") or "") for column in columns
            ]
            body = [
                [str(row.get("label") or "")]
                + [
                    str((row.get("cells") or {}).get(str(column.get("key")), ""))
                    for column in columns
                ]
                for row in rows
            ]
        else:
            header = [str(item) for item in columns]
            body = [
                [
                    str(row.get("label") or ""),
                    str(row.get("value") or ""),
                    str(row.get("missing") or ""),
                ]
                for row in rows
            ]
        if header:
            _add_table(doc, header, body)
        for note in rendered.get("notes") or []:
            _add_note(doc, str(note))


def _section_references(doc: Document, analyses: list[dict[str, Any]]) -> None:
    references: list[str] = []
    for analysis in analyses:
        for reference in (analysis.get("result") or {}).get("references") or []:
            text = str(reference).strip()
            if text and text not in references:
                references.append(text)
    if not references:
        return
    doc.add_heading("6. Methods references", level=1)
    for reference in references:
        doc.add_paragraph(reference, style="List Bullet")


def _p(value: Any) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return "–"
    if number != number:  # NaN
        return "–"
    return "<0.001" if number < 0.001 else f"{number:.3f}"


def build_dossier(
    *,
    table: dict[str, Any],
    project: dict[str, Any],
    versions: list[dict[str, Any]],
    population: dict[str, Any] | None,
    flow: dict[str, Any] | None,
    quality: dict[str, Any] | None,
    steps: list[dict[str, Any]],
    analyses: list[dict[str, Any]],
    generated_at: str,
) -> bytes:
    """Build the dossier and return the ``.docx`` bytes."""
    document = Document()
    title = document.add_heading("Study dossier", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle = document.add_paragraph()
    run = subtitle.add_run(
        f"{project.get('name') or project.get('id')} — {table.get('name')}"
    )
    run.font.size = Pt(12)
    warning = document.add_paragraph()
    run = warning.add_run(_DISCLAIMER)
    run.bold = True
    run.font.size = Pt(9)

    _section_provenance(document, table, project, versions, population)
    _section_flow(document, flow)
    _section_handling(document, steps, quality)
    _section_missing(document, analyses)
    _section_results(document, analyses)
    _section_references(document, analyses)

    footer = document.add_paragraph()
    run = footer.add_run(f"Generated {generated_at} by the Gazzali cohort module.")
    run.font.size = Pt(8)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
